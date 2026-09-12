"""
Comprehensive test suite covering Document Structure Understanding Scenarios A through Z
plus multi-format extraction tests (PDF, TXT, Markdown, EPUB, DOCX).
"""

from pathlib import Path
import docx
import ebooklib
from ebooklib import epub
import fitz
import pytest

from app.services.document.engine import DocumentStructureEngine
from app.services.document.extractors import (
    DOCXExtractor,
    EPUBExtractor,
    MarkdownExtractor,
    PDFExtractor,
    TXTExtractor,
)
from app.services.document.heading_detector import HeadingCandidate
from app.services.document.models import NodeType
from app.services.document.toc import ParsedTOC, TOCEntry, reconcile_toc_with_candidates
from app.services.document.validator import validate_document_tree


def test_scenario_a_flat_30_chapters():
    """Scenario A: 30 chapters, no parts."""
    text_blocks = []
    for i in range(1, 31):
        text_blocks.append(f"CHAPTER {i}\nThis is the content of chapter {i}. It describes the events in detail.")
    full_text = "\n\n".join(text_blocks)

    tree = DocumentStructureEngine.parse_text_stream(full_text, "Book of 30 Chapters")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 30
    assert chapters[0].title == "Chapter 1"
    assert chapters[29].title == "Chapter 30"
    assert "events in detail" in chapters[29].text


def test_scenario_b_parts_and_chapters():
    """Scenario B: 5 parts × 6 chapters."""
    text_blocks = []
    for p in range(1, 6):
        text_blocks.append(f"PART {p}\nThis is the opening of part {p}.")
        for c in range(1, 7):
            text_blocks.append(f"Chapter {c}\nContent for part {p} chapter {c}.")
    full_text = "\n\n".join(text_blocks)

    tree = DocumentStructureEngine.parse_text_stream(full_text, "5 Parts Book")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 30
    assert "Part 1" in chapters[0].title or "Part 1 - Chapter 1" in chapters[1].title


def test_scenario_c_parts_with_sections_no_chapters():
    """Scenario C: Parts containing sections but no chapters."""
    raw = (
        "PART ONE\n\n"
        "1.1 First Principles\n"
        "Here are the first principles of the system.\n\n"
        "1.2 Secondary Rules\n"
        "Here are secondary rules.\n\n"
        "PART TWO\n\n"
        "2.1 Execution\n"
        "Execution details go here."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Manual")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 3
    assert any("First Principles" in ch.title for ch in chapters)
    assert any("Execution" in ch.title for ch in chapters)


def test_scenario_d_nested_book_part_chapter_section():
    """Scenario D: Book → Part → Chapter → Section."""
    raw = (
        "BOOK I\n\n"
        "PART ONE\n\n"
        "CHAPTER 1\n\n"
        "1.1 Foundations\n"
        "Foundational material is outlined here.\n\n"
        "1.2 Frameworks\n"
        "Frameworks are explained here."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Comprehensive Guide")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 2
    assert "Foundations" in chapters[0].title
    assert "Frameworks" in chapters[1].title


def test_scenario_e_prologue_chapters_epilogue():
    """Scenario E: Prologue + chapters + epilogue."""
    raw = (
        "PROLOGUE\n"
        "Long before the journey began, the world was silent.\n\n"
        "CHAPTER 1\n"
        "The first step of the journey.\n\n"
        "CHAPTER 2\n"
        "The second step of the journey.\n\n"
        "EPILOGUE\n"
        "And so the journey concluded peacefully."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Journey")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 4
    assert chapters[0].title == "Prologue"
    assert chapters[1].title == "Chapter 1"
    assert chapters[2].title == "Chapter 2"
    assert chapters[3].title == "Epilogue"


def test_scenario_f_preface_intro_chapters_appendix():
    """Scenario F: Preface + introduction + chapters + appendix."""
    raw = (
        "PREFACE\n"
        "Why I wrote this book.\n\n"
        "INTRODUCTION\n"
        "How to read this book.\n\n"
        "CHAPTER 1\n"
        "The main concept.\n\n"
        "APPENDIX A\n"
        "Reference tables and formulas."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Textbook")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 4
    assert chapters[0].title == "Preface"
    assert chapters[1].title == "Introduction"
    assert chapters[2].title == "Chapter 1"
    assert "Appendix" in chapters[3].title


def test_scenario_g_roman_numeral_chapters():
    """Scenario G: Roman numeral chapters."""
    raw = (
        "CHAPTER I\n"
        "Roman chapter one.\n\n"
        "CHAPTER IV\n"
        "Roman chapter four.\n\n"
        "CHAPTER XII\n"
        "Roman chapter twelve."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Roman Book")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 3
    assert chapters[0].title == "Chapter I"
    assert chapters[1].title == "Chapter IV"
    assert chapters[2].title == "Chapter XII"


def test_scenario_h_numbered_chapters_without_word_chapter():
    """Scenario H: Numbered chapters without the word 'chapter'."""
    raw = (
        "1. The Awakening\n"
        "He woke up in the morning light.\n\n"
        "2. The Journey North\n"
        "They traveled north across the plains.\n\n"
        "3. The City Gates\n"
        "They finally reached the city."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Adventure")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 3
    assert "1 The Awakening" in chapters[0].title or "1. The Awakening" in chapters[0].title
    assert "2 The Journey North" in chapters[1].title or "2. The Journey North" in chapters[1].title


def test_scenario_i_no_explicit_headings():
    """Scenario I: Books with no explicit headings."""
    raw = (
        "Paragraph one of an unheaded essay.\n\n"
        "Paragraph two continues the discussion.\n\n"
        "Paragraph three adds another point.\n\n"
        "Paragraph four elaborates further.\n\n"
        "Paragraph five discusses implications.\n\n"
        "Paragraph six offers counterpoints.\n\n"
        "Paragraph seven evaluates evidence.\n\n"
        "Paragraph eight concludes the essay."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Untitled Essay")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 1
    assert "Paragraph one" in chapters[0].text
    assert "concludes the essay" in chapters[-1].text


def test_scenario_j_scanned_or_single_block_text():
    """Scenario J: Scanned / flat text."""
    raw = "A continuous single block of text without blank lines or standard chapter headers."
    tree = DocumentStructureEngine.parse_text_stream(raw, "Single Block")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 1
    assert "continuous single block" in chapters[0].text


def test_scenario_k_toc_detection():
    """Scenario K: Printed Table of Contents page."""
    raw = (
        "TABLE OF CONTENTS\n"
        "Chapter 1: The Start .......... 1\n"
        "Chapter 2: The Middle ......... 5\n"
        "Chapter 3: The End ............ 10\n\n"
        "CHAPTER 1: The Start\n"
        "This is the start.\n\n"
        "CHAPTER 2: The Middle\n"
        "This is the middle.\n\n"
        "CHAPTER 3: The End\n"
        "This is the end."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Book With TOC")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 3
    assert any("The Start" in c.title for c in chapters)
    assert any("The End" in c.title for c in chapters)


def test_scenario_l_no_toc():
    """Scenario L: Book without TOC."""
    raw = "CHAPTER 1\nBody text one.\n\nCHAPTER 2\nBody text two."
    tree = DocumentStructureEngine.parse_text_stream(raw, "No TOC Book")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 2


def test_scenario_m_nested_multilevel():
    """Scenario M: Multiple nested structural levels."""
    raw = (
        "PART ONE\n\n"
        "CHAPTER 1\n\n"
        "1.1 Foundations\n"
        "Text in 1.1\n\n"
        "1.2 Framework\n"
        "Text in 1.2\n\n"
        "CHAPTER 2\n\n"
        "2.1 Implementation\n"
        "Text in 2.1"
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Multi Level")
    report = validate_document_tree(tree, raw_text=raw)

    assert report.is_valid
    assert len(tree.root.children) >= 1


def test_scenario_n_running_headers_footers_rejection():
    """Scenario N: Running headers/footers with page numbers not treated as chapters."""
    from app.services.document.layout import LayoutBlock, PageLayout, _filter_headers_and_footers

    pages = []
    for p in range(1, 5):
        blocks = [
            LayoutBlock(text="ORWELL | 1984", spans=[], bbox=(100, 20, 300, 40), page_num=p, avg_font_size=8.0, is_bold=False, is_all_caps=True, line_count=1),
            LayoutBlock(text=f"CHAPTER {p}\nThis is body text on page {p}.", spans=[], bbox=(100, 100, 500, 400), page_num=p, avg_font_size=10.0, is_bold=False, is_all_caps=False, line_count=2),
            LayoutBlock(text=str(p), spans=[], bbox=(250, 750, 270, 770), page_num=p, avg_font_size=8.0, is_bold=False, is_all_caps=False, line_count=1),
        ]
        pages.append(PageLayout(page_num=p, width=612, height=792, blocks=blocks, raw_text=""))

    _filter_headers_and_footers(pages)

    for page in pages:
        texts = [b.text for b in page.blocks]
        assert "ORWELL | 1984" not in texts
        assert str(page.page_num) not in texts


def test_scenario_o_multi_line_headings():
    """Scenario O: Chapter titles spanning multiple lines (e.g. CHAPTER 1 and THE JOURNEY)."""
    raw = (
        "CHAPTER 1\n"
        "THE BOY WHO LIVED\n\n"
        "Mr. and Mrs. Dursley, of number four, Privet Drive, were proud to say that they were perfectly normal.\n\n"
        "CHAPTER 2\n"
        "THE VANISHING GLASS\n\n"
        "Nearly ten years had passed since the Dursleys had woken up to find their nephew on the front doorstep."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Wizard Novel")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 2
    assert "Boy Who Lived" in chapters[0].title or "The Boy Who Lived" in chapters[0].title
    assert "Vanishing Glass" in chapters[1].title or "The Vanishing Glass" in chapters[1].title


def test_scenario_p_chapter_spans_pages():
    """Scenario P: Chapter begins and continues across text sections."""
    raw = (
        "CHAPTER 1\n"
        "First sentence of chapter 1.\n\n"
        "Second paragraph continuing chapter 1 onto next section.\n\n"
        "CHAPTER 2\n"
        "First sentence of chapter 2."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Two Pages")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 2
    assert "First sentence of chapter 1" in chapters[0].text
    assert "Second paragraph continuing" in chapters[0].text


def test_scenario_q_dialogue_quotes_not_split():
    """Scenario Q: Chapter containing internal dialogue and quotes that must NOT split."""
    raw = (
        "CHAPTER 1\n"
        "Winston was sitting in his cubicle.\n\n"
        '"Chapter 1 of the manual is very clear," O\'Brien said.\n\n'
        "He opened Goldstein's book to Chapter 1. IGNORANCE IS STRENGTH.\n\n"
        "Throughout recorded time, there have been three kinds of people.\n\n"
        '"Part Two will come tomorrow," whispered Julia.'
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Dystopia")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 1
    assert "IGNORANCE IS STRENGTH" in chapters[0].text
    assert "whispered Julia" in chapters[0].text


def test_scenario_r_hindi_headings():
    """Scenario R: Hindi structural headings (भाग, अध्याय, प्रस्तावना, उपसंहार)."""
    raw = (
        "प्रस्तावना\n"
        "यह पुस्तक का आरंभिक परिचय है।\n\n"
        "भाग १\n"
        "पहला भाग यहाँ से शुरू होता है।\n\n"
        "अध्याय १\n"
        "यह पहले अध्याय की मुख्य सामग्री है।\n\n"
        "अध्याय २\n"
        "यह दूसरे अध्याय की सामग्री है।\n\n"
        "उपसंहार\n"
        "यह पुस्तक का अंतिम उपसंहार है।"
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Hindi Book")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 4
    assert any("Prologue" in c.title or "प्रस्तावना" in c.title for c in chapters)
    assert any("Chapter" in c.title or "अध्याय" in c.title for c in chapters)
    assert any("Epilogue" in c.title or "उपसंहार" in c.title for c in chapters)



def test_scenario_s_mixed_english_hindi():
    """Scenario S: Mixed English and Hindi document."""
    raw = (
        "PREFACE\n"
        "Welcome to the bilingual edition.\n\n"
        "अध्याय १: The Journey\n"
        "यात्रा यहाँ से आरंभ होती है। It was a sunny morning.\n\n"
        "CHAPTER 2: निष्कर्ष\n"
        "The conclusion of the bilingual journey."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Bilingual Work")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 3
    assert "Preface" in chapters[0].title
    assert "The Journey" in chapters[1].title or "अध्याय" in chapters[1].title
    assert "निष्कर्ष" in chapters[2].title or "Chapter 2" in chapters[2].title


def test_scenario_t_multicolumn_pdf_layout():
    """Scenario T: Multi-column PDF layout blocks correctly ordered by top-to-bottom reading order."""
    from app.services.document.layout import LayoutBlock

    col1 = LayoutBlock(text="Left column text top", bbox=(50, 100, 250, 200), page_num=1, avg_font_size=10.0, is_bold=False, is_all_caps=False, line_count=5)
    col2 = LayoutBlock(text="Right column text top", bbox=(300, 100, 500, 200), page_num=1, avg_font_size=10.0, is_bold=False, is_all_caps=False, line_count=5)
    header = LayoutBlock(text="CHAPTER 1", bbox=(50, 40, 500, 70), page_num=1, avg_font_size=16.0, is_bold=True, is_all_caps=True, line_count=1)

    blocks = [col1, col2, header]
    blocks.sort(key=lambda blk: (round(blk.bbox[1] / 10.0) * 10.0, blk.bbox[0]))

    assert blocks[0].text == "CHAPTER 1"
    assert blocks[1].text == "Left column text top"
    assert blocks[2].text == "Right column text top"


def test_scenario_u_repeated_body_typography_not_split():
    """Scenario U: Repeated bold words in body text (e.g. Note:, Warning:) do not create false chapters."""
    raw = (
        "CHAPTER 1\n"
        "Here is the main text.\n\n"
        "Important Note:\n"
        "Make sure to follow safety instructions.\n\n"
        "Another Note:\n"
        "Review the checklist carefully.\n\n"
        "CHAPTER 2\n"
        "Here is chapter two."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Manual with Notes")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 2
    assert "Important Note" in chapters[0].text
    assert "Another Note" in chapters[0].text


def test_scenario_v_toc_page_offset_reconciliation():
    """Scenario V: TOC page-number offset reconciliation."""
    toc = ParsedTOC(
        entries=[
            TOCEntry(level=2, title="Chapter 1: Beginning", page=1),
            TOCEntry(level=2, title="Chapter 2: Middle", page=15),
        ]
    )
    candidates = [
        HeadingCandidate(
            title="Chapter 1 - Beginning",
            clean_title="Beginning",
            node_type=NodeType.CHAPTER,
            level=2,
            ordinal=1,
            page_num=5,
            raw_block_text="CHAPTER 1: Beginning",
            confidence=0.95,
            detection_method="lexical_chapter",
        ),
        HeadingCandidate(
            title="Chapter 2 - Middle",
            clean_title="Middle",
            node_type=NodeType.CHAPTER,
            level=2,
            ordinal=2,
            page_num=19,
            raw_block_text="CHAPTER 2: Middle",
            confidence=0.95,
            detection_method="lexical_chapter",
        ),
    ]

    report = reconcile_toc_with_candidates(toc, candidates, total_pages=50)
    assert report["reconciled"] is True
    assert report["page_offset"] == 4
    assert report["matches"] == 2


def test_scenario_w_toc_body_disagreement():
    """Scenario W: Disagreement between TOC and body headings is recorded without crashing."""
    toc = ParsedTOC(
        entries=[
            TOCEntry(level=2, title="Chapter 1: Found", page=1),
            TOCEntry(level=2, title="Chapter 2: Missing From Body", page=20),
        ]
    )
    candidates = [
        HeadingCandidate(
            title="Chapter 1 - Found",
            clean_title="Found",
            node_type=NodeType.CHAPTER,
            level=2,
            ordinal=1,
            page_num=1,
            raw_block_text="CHAPTER 1: Found",
            confidence=0.95,
            detection_method="lexical_chapter",
        )
    ]

    report = reconcile_toc_with_candidates(toc, candidates, total_pages=30)
    assert report["matches"] == 1
    assert len(report["disagreements"]) == 1
    assert "Missing From Body" in report["disagreements"][0]


def test_scenario_x_document_only_body_text():
    """Scenario X: A document containing only body text produces a valid single chapter with zero loss."""
    raw = "Just a single paragraph of plain body text without any headings or structural indicators."
    tree = DocumentStructureEngine.parse_text_stream(raw, "Single Essay")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 1
    assert "Just a single paragraph" in chapters[0].text
    report = validate_document_tree(tree, raw_text=raw)
    assert report.is_valid is True
    assert report.text_loss_rate == 0.0


def test_scenario_y_appendix_glossary_bibliography():
    """Scenario Y: Appendix, Glossary, and Bibliography following chapters."""
    raw = (
        "CHAPTER 1\n"
        "The story begins here.\n\n"
        "CHAPTER 2\n"
        "The story ends here.\n\n"
        "APPENDIX A\n"
        "Supplementary derivations.\n\n"
        "GLOSSARY\n"
        "Definitions of key terminology.\n\n"
        "BIBLIOGRAPHY\n"
        "List of published references."
    )
    tree = DocumentStructureEngine.parse_text_stream(raw, "Academic Book")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) == 5
    assert chapters[0].title == "Chapter 1"
    assert chapters[1].title == "Chapter 2"
    assert "Appendix" in chapters[2].title
    assert "Glossary" in chapters[3].title
    assert "Bibliography" in chapters[4].title


def test_scenario_z_nested_sections_without_conventional_numbering():
    """Scenario Z: Nested sections with descriptive titles without conventional numbering."""
    raw = (
        "# System Architecture\n\n"
        "Overview of the system.\n\n"
        "## Ingestion Layer\n\n"
        "Details of the ingestion pipeline.\n\n"
        "### PDF Parser\n\n"
        "PyMuPDF structured extraction.\n\n"
        "## Audio Synthesis Layer\n\n"
        "Kokoro TTS generation."
    )
    tree = DocumentStructureEngine.parse_markdown(raw, "Architecture Spec")
    chapters = tree.to_narratable_chapters()

    assert len(chapters) >= 3
    assert any("Ingestion Layer" in c.title for c in chapters)
    assert any("Audio Synthesis Layer" in c.title for c in chapters)


# --- Multi-Format Ingestion Tests ---

def test_multiformat_txt_extraction(tmp_path: Path):
    p = tmp_path / "sample.txt"
    p.write_text("CHAPTER 1\n\nHello from plain text.\n\nCHAPTER 2\n\nSecond chapter.", encoding="utf-8")
    extractor = TXTExtractor()
    assert extractor.can_handle(p) is True

    norm_doc = extractor.extract(p)
    assert norm_doc.source_type == "txt"
    assert len(norm_doc.pages) == 1
    assert "Hello from plain text." in norm_doc.raw_text


def test_multiformat_markdown_extraction(tmp_path: Path):
    p = tmp_path / "sample.md"
    p.write_text("# Chapter 1: Foundations\n\nContent here.\n\n## Section 1.1\n\nSubcontent.", encoding="utf-8")
    extractor = MarkdownExtractor()
    assert extractor.can_handle(p) is True

    norm_doc = extractor.extract(p)
    assert norm_doc.source_type == "markdown"
    assert len(norm_doc.toc_entries) >= 2
    assert norm_doc.toc_entries[0].title == "Chapter 1: Foundations"


def test_multiformat_epub_extraction(tmp_path: Path):
    p = tmp_path / "test.epub"
    book = epub.EpubBook()
    book.set_title("Test EPUB Book")
    book.set_language("en")
    book.add_author("Novela Author")

    c1 = epub.EpubHtml(title="Chapter 1", file_name="chap_1.xhtml", lang="en")
    c1.content = "<h1>Chapter 1: Beginning</h1><p>First paragraph of EPUB text.</p>"
    book.add_item(c1)

    # Add standard NCX and Nav items so reader does not crash on missing TOC navigation
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", c1]
    book.toc = [c1]
    epub.write_epub(str(p), book)

    extractor = EPUBExtractor()
    assert extractor.can_handle(p) is True

    norm_doc = extractor.extract(p)
    assert norm_doc.source_type == "epub"
    assert norm_doc.title == "Test EPUB Book"
    assert "First paragraph of EPUB text." in norm_doc.raw_text


def test_multiformat_docx_extraction(tmp_path: Path):
    p = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_heading("Chapter 1: The Beginning", level=1)
    doc.add_paragraph("This is the first paragraph of the DOCX test document.")
    doc.add_heading("Chapter 2: The Next Step", level=1)
    doc.add_paragraph("This is the second paragraph.")
    doc.save(str(p))

    extractor = DOCXExtractor()
    assert extractor.can_handle(p) is True

    norm_doc = extractor.extract(p)
    assert norm_doc.source_type == "docx"
    assert len(norm_doc.toc_entries) == 2
    assert "first paragraph" in norm_doc.raw_text


def test_multiformat_pdf_extraction(tmp_path: Path):
    p = tmp_path / "test.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "CHAPTER 1\n\nThis is a PyMuPDF synthesized test PDF.", fontsize=14)
    doc.save(str(p))
    doc.close()

    extractor = PDFExtractor()
    assert extractor.can_handle(p) is True

    norm_doc = extractor.extract(p)
    assert norm_doc.source_type == "pdf"
    assert len(norm_doc.pages) == 1
    assert "PyMuPDF synthesized" in norm_doc.raw_text
